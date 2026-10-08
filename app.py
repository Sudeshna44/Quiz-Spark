import json
import os
import re
from io import BytesIO
from typing import Any

import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI

load_dotenv()


@st.cache_resource
def get_openrouter_client() -> OpenAI:
    """Create a single cached client for the OpenRouter API."""
    return OpenAI(
        api_key=OPENROUTER_API_KEY,
        base_url=OPENROUTER_BASE_URL,
        timeout=600.0,
    )


@st.cache_data(show_spinner=False)
def build_word_document_cached(
    question_payload: str,
    *,
    grade: str,
    subject: str,
    topic: str,
    difficulty: str,
    question_type: str,
) -> bytes:
    """Generate and cache the Word download for a specific quiz payload."""
    questions = json.loads(question_payload)
    return make_word_document(
        questions,
        grade=grade,
        subject=subject,
        topic=topic,
        difficulty=difficulty,
        question_type=question_type,
    )
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openai import APIConnectionError, APIError, APIStatusError, OpenAI


# ============================================================
# OPENROUTER CONFIGURATION
# ============================================================

OPENROUTER_BASE_URL = os.getenv(
    "OPENROUTER_BASE_URL",
    "https://openrouter.ai/api/v1",
)
OPENROUTER_MODEL = os.getenv("OPENROUTER_MODEL", "openrouter/free")
OPENROUTER_API_KEY = (
    os.getenv("OPENROUTER_API_KEY")
    or os.getenv("OLLAMA_API_KEY", "")
)
OPENROUTER_MAX_TOKENS = int(os.getenv("OPENROUTER_MAX_TOKENS", "1200"))
OPENROUTER_TEMPERATURE = float(os.getenv("OPENROUTER_TEMPERATURE", "0.3"))


# ============================================================
# QUESTION TYPES
# ============================================================

QUESTION_TYPES = (
    "Multiple choice",
    "True or false",
    "Short answer",
)

TYPE_TO_KEY = {
    "Multiple choice": "multiple_choice",
    "True or false": "true_false",
    "Short answer": "short_answer",
}


# ============================================================
# CLEAN JSON RESPONSE
# ============================================================

def clean_json_response(content: str) -> Any:
    """
    Parse JSON returned by the model.

    Handles both:
    1. Normal JSON
    2. JSON surrounded by ```json ... ```
    """

    cleaned = content.strip()

    # Remove markdown code fences if the model returns them
    fenced_match = re.fullmatch(
        r"```(?:json)?\s*(.*?)\s*```",
        cleaned,
        re.DOTALL | re.IGNORECASE,
    )

    if fenced_match:
        cleaned = fenced_match.group(1).strip()

    return json.loads(cleaned)


# ============================================================
# VALIDATE QUIZ
# ============================================================

def validate_quiz(
    data: Any,
    expected_count: int,
    question_type: str,
) -> list[dict[str, Any]]:

    if not isinstance(data, dict):
        raise ValueError("The model response was not a JSON object.")

    if not isinstance(data.get("questions"), list):
        raise ValueError("The model response did not include a questions list.")

    questions = data["questions"]

    if len(questions) != expected_count:
        raise ValueError(
            f"The model returned {len(questions)} questions "
            f"instead of {expected_count}. Please try again."
        )

    expected_type = TYPE_TO_KEY[question_type]

    for number, question in enumerate(questions, start=1):

        if not isinstance(question, dict):
            raise ValueError(
                f"Question {number} is not in the expected format."
            )

        # Question text
        if (
            not isinstance(question.get("question"), str)
            or not question["question"].strip()
        ):
            raise ValueError(
                f"Question {number} is missing its question text."
            )

        # Correct answer
        if (
            not isinstance(question.get("correct_answer"), str)
            or not question["correct_answer"].strip()
        ):
            raise ValueError(
                f"Question {number} is missing its answer."
            )

        # Explanation
        if not isinstance(question.get("explanation"), str):
            raise ValueError(
                f"Question {number} is missing its explanation."
            )

        # Multiple choice validation
        if expected_type == "multiple_choice":

            options = question.get("options")

            if (
                not isinstance(options, list)
                or len(options) != 4
                or not all(
                    isinstance(option, str) and option.strip()
                    for option in options
                )
                or question["correct_answer"] not in options
            ):
                raise ValueError(
                    f"Question {number} has invalid multiple-choice "
                    f"options or answer."
                )

            question["options"] = options

        # True / False validation
        elif expected_type == "true_false":

            answer = question["correct_answer"].strip().lower()

            if answer not in ("true", "false"):
                raise ValueError(
                    f"Question {number}'s true-or-false answer "
                    f"must be True or False."
                )

            question["correct_answer"] = answer.capitalize()

        # Set internal question type
        question["type"] = expected_type

    return questions


# ============================================================
# CREATE QUIZ USING OPENROUTER
# ============================================================

@st.cache_data(show_spinner=False, ttl=3600)
def make_quiz(
    *,
    grade: str,
    subject: str,
    topic: str,
    difficulty: str,
    count: int,
    question_type: str,
) -> list[dict[str, Any]]:

    # --------------------------------------------------------
    # Reuse a cached OpenAI-compatible client connected to OpenRouter
    # --------------------------------------------------------

    if not OPENROUTER_API_KEY:
        raise ValueError(
            "Set OPENROUTER_API_KEY in your .env file to generate quizzes."
        )

    client = get_openrouter_client()

    kind = TYPE_TO_KEY[question_type]

    # --------------------------------------------------------
    # Question-specific instructions
    # --------------------------------------------------------

    if kind == "multiple_choice":

        options_instruction = """
For every question:
- Include exactly four options.
- options must contain four strings.
- correct_answer must exactly match one of the four options.
"""

        answer_instruction = """
For multiple choice questions, correct_answer must exactly
match one of the options.
"""

    elif kind == "true_false":

        options_instruction = """
For every question:
- Set options to an empty array.
- correct_answer must be either True or False.
"""

        answer_instruction = """
For true or false questions, correct_answer must be exactly
True or False.
"""

    else:

        options_instruction = """
For every question:
- Set options to an empty array.
- correct_answer should be a short and accurate answer.
"""

        answer_instruction = """
For short answer questions, correct_answer should be a concise
model answer.
"""

    # --------------------------------------------------------
    # Prompt
    # --------------------------------------------------------

    prompt = f"""
Create exactly {count} original, accurate and age-appropriate
{kind} quiz questions.

Student standard/grade: {grade}

Subject: {subject}

Topic: {topic}

Difficulty: {difficulty}

Use joyful, encouraging and classroom-friendly wording.

Keep the questions:
- Clear
- Educational
- Age appropriate
- Accurate
- Engaging

Avoid trick questions.

Include a short helpful explanation for every answer.

{options_instruction}

{answer_instruction}

IMPORTANT:
Return ONLY valid JSON.
Do not include markdown.
Do not include ```json.
Do not include any explanation outside the JSON.

Return exactly this structure:

{{
    "questions": [
        {{
            "question": "...",
            "options": [],
            "correct_answer": "...",
            "explanation": "..."
        }}
    ]
}}
"""

    # --------------------------------------------------------
    # Call OpenRouter
    # --------------------------------------------------------

    response = client.chat.completions.create(
        model=OPENROUTER_MODEL,
        messages=[
            {
                "role": "system",
                "content": (
                    "You are a warm, accurate teacher who creates "
                    "fun learning activities. Keep responses short, "
                    "clear, and valid JSON only."
                ),
            },
            {
                "role": "user",
                "content": prompt,
            },
        ],
        temperature=OPENROUTER_TEMPERATURE,
        max_tokens=OPENROUTER_MAX_TOKENS,
    )

    # --------------------------------------------------------
    # Read response
    # --------------------------------------------------------

    content = response.choices[0].message.content

    if not content:
        raise ValueError(
            "OpenRouter returned an empty response. "
            "Please check the model and try again."
        )

    # --------------------------------------------------------
    # Parse JSON
    # --------------------------------------------------------

    try:
        quiz_data = clean_json_response(content)

    except json.JSONDecodeError as exc:

        raise ValueError(
            "OpenRouter returned a response that was not valid JSON.\n\n"
            f"Raw response:\n{content}"
        ) from exc

    # --------------------------------------------------------
    # Validate
    # --------------------------------------------------------

    return validate_quiz(
        quiz_data,
        count,
        question_type,
    )


# ============================================================
# CREATE WORD DOCUMENT
# ============================================================

def make_word_document(
    questions: list[dict[str, Any]],
    *,
    grade: str,
    subject: str,
    topic: str,
    difficulty: str,
    question_type: str,
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # --------------------------------------------------------
    # Normal font
    # --------------------------------------------------------

    normal = document.styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(46, 58, 89)

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    title = document.add_heading(
        "Your Bright Ideas Quiz!",
        0,
    )

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title.runs[0].font.color.rgb = RGBColor(
        89,
        75,
        174,
    )

    subtitle = document.add_paragraph(
        "You've got this! Take your time and enjoy the challenge."
    )

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --------------------------------------------------------
    # Quiz details
    # --------------------------------------------------------

    details = document.add_paragraph()

    details.add_run("Standard: ").bold = True
    details.add_run(f"{grade}    ")

    details.add_run("Subject: ").bold = True
    details.add_run(f"{subject}\n")

    details.add_run("Topic: ").bold = True
    details.add_run(f"{topic}    ")

    details.add_run("Difficulty: ").bold = True
    details.add_run(difficulty)

    document.add_paragraph(
        "Name: ________________________________    "
        "Date: ________________"
    )

    document.add_paragraph(
        f"Instructions: Read each {question_type.lower()} "
        "question carefully. Have fun learning!"
    )

    # --------------------------------------------------------
    # Questions
    # --------------------------------------------------------

    for number, question in enumerate(
        questions,
        start=1,
    ):

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(10)

        paragraph.add_run(
            f"{number}. {question['question']}"
        ).bold = True

        # Multiple choice
        if question["type"] == "multiple_choice":

            for index, option in enumerate(
                question["options"]
            ):

                document.add_paragraph(
                    f"   {chr(65 + index)}. {option}"
                )

        # True / False
        elif question["type"] == "true_false":

            document.add_paragraph(
                "   ☐ True       ☐ False"
            )

        # Short answer
        else:

            document.add_paragraph(
                "   ________________________________________________"
            )

            document.add_paragraph(
                "   ________________________________________________"
            )

    # --------------------------------------------------------
    # Answer key
    # --------------------------------------------------------

    document.add_page_break()

    answer_title = document.add_heading(
        "Answer Key & Learning Sparks",
        0,
    )

    answer_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(
        "Great effort! Use these answers to check your thinking "
        "and learn something new."
    )

    for number, question in enumerate(
        questions,
        start=1,
    ):

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(8)

        paragraph.add_run(
            f"{number}. {question['correct_answer']}"
        ).bold = True

        document.add_paragraph(
            f"Why? {question['explanation']}"
        )

    # --------------------------------------------------------
    # Convert document to bytes
    # --------------------------------------------------------

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


# ============================================================
# STREAMLIT PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Quiz Spark | Joyful Quiz Maker",
    page_icon="✨",
    layout="centered",
)


# ============================================================
# PAGE UI
# ============================================================

st.title("✨ Quiz Spark")

st.subheader(
    "Turn today's topic into a joyful little challenge!"
)

st.write(
    "Tell us what your class is learning, and we'll make "
    "a ready-to-print quiz in Word."
)


# ============================================================
# QUIZ FORM
# ============================================================

with st.form("quiz_form"):

    grade = st.selectbox(
        "Student standard / grade",
        [f"Grade {grade}" for grade in range(1, 13)],
    )

    subject = st.text_input(
        "Subject",
        placeholder="e.g. Science, Mathematics, English",
    )

    topic = st.text_input(
        "Topic to explore",
        placeholder="e.g. The water cycle",
    )

    left, right = st.columns(2)

    with left:

        difficulty = st.selectbox(
            "Challenge level",
            [
                "Gentle warm-up",
                "Just right",
                "Brain booster",
            ],
        )

        question_type = st.selectbox(
            "Question style",
            QUESTION_TYPES,
        )

    with right:

        count = st.slider(
            "Number of questions",
            min_value=3,
            max_value=15,
            value=5,
        )

    generate = st.form_submit_button(
        "✨ Make my quiz!",
        type="primary",
        use_container_width=True,
    )


# ============================================================
# GENERATE QUIZ
# ============================================================

if generate:

    if not subject.strip() or not topic.strip():

        st.error(
            "Please add both a subject and a topic "
            "so we can make your quiz just right."
        )

    else:

        with st.spinner(
            "✨ Your quiz is getting its thinking cap on..."
        ):

            try:

                questions = make_quiz(
                    grade=grade,
                    subject=subject.strip(),
                    topic=topic.strip(),
                    difficulty=difficulty,
                    count=count,
                    question_type=question_type,
                )

            # ------------------------------------------------
            # Connection error
            # ------------------------------------------------

            except APIConnectionError as exc:

                st.error(
                    "❌ Could not connect to OpenRouter."
                )

                st.code(
                    str(exc),
                    language="text",
                )

                st.info(
                    "Check your internet connection and the "
                    "OPENROUTER_BASE_URL setting."
                )

            # ------------------------------------------------
            # HTTP/API error
            # ------------------------------------------------

            except APIStatusError as exc:

                st.error(
                    f"❌ OpenRouter API error "
                    f"(HTTP {exc.status_code})"
                )

                st.code(
                    str(exc),
                    language="text",
                )

            # ------------------------------------------------
            # Other OpenAI API error
            # ------------------------------------------------

            except APIError as exc:

                st.error(
                    "❌ OpenRouter API request failed."
                )

                st.code(
                    str(exc),
                    language="text",
                )

            # ------------------------------------------------
            # JSON / validation error
            # ------------------------------------------------

            except ValueError as exc:

                st.error(
                    "❌ Quiz generation error"
                )

                st.code(
                    str(exc),
                    language="text",
                )

            # ------------------------------------------------
            # Unexpected error
            # ------------------------------------------------

            except Exception as exc:

                st.error(
                    "❌ Unexpected error"
                )

                st.code(
                    repr(exc),
                    language="text",
                )

            # ------------------------------------------------
            # Successful generation
            # ------------------------------------------------

            else:

                st.session_state["quiz"] = {
                    "questions": questions,
                    "grade": grade,
                    "subject": subject.strip(),
                    "topic": topic.strip(),
                    "difficulty": difficulty,
                    "question_type": question_type,
                }


# ============================================================
# DISPLAY GENERATED QUIZ
# ============================================================

if quiz := st.session_state.get("quiz"):

    st.success(
        "🎉 Your quiz is ready! "
        "You can preview it here or download the Word document."
    )

    # --------------------------------------------------------
    # Preview questions
    # --------------------------------------------------------

    for number, question in enumerate(
        quiz["questions"],
        start=1,
    ):

        with st.expander(
            f"Question {number}: {question['question']}",
            expanded=True,
        ):

            if question["type"] == "multiple_choice":

                for option in question["options"]:

                    st.write(
                        f"• {option}"
                    )

            elif question["type"] == "true_false":

                st.write("True / False")

            st.caption(
                f"Answer: {question['correct_answer']}"
            )

            st.write(
                question["explanation"]
            )

    # --------------------------------------------------------
    # Create Word document
    # --------------------------------------------------------

    word_file = build_word_document_cached(
        json.dumps(quiz["questions"]),
        grade=quiz["grade"],
        subject=quiz["subject"],
        topic=quiz["topic"],
        difficulty=quiz["difficulty"],
        question_type=quiz["question_type"],
    )

    # --------------------------------------------------------
    # Download button
    # --------------------------------------------------------

    st.download_button(
        label="📄 Download quiz as a Word document",
        data=word_file,
        file_name="quiz_spark.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        type="primary",
        use_container_width=True,
    )